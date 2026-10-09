# -*- coding: utf-8 -*-
"""适配 BI 新版导出(2026-09-30 起)。

格式变更:
  旧:sheet「官网监控明细数据」 维度 = 日期 × 付费状态 × D0 × 国家   (85列)
  新:sheet「官网大盘数据」     维度 = 日期 × 渠道(SEO/fb/未知)      (86列)
      指标改名:触达付费集uv → 快捷支付面板曝光uv;创建订单uv → 商品点击uv
      新增:快捷支付面板曝光率、商品展示uv、商品点击率

后果:**国家 / 付费状态 / D0 三个维度在日报里没有了**,所以:
  - 大盘(收入/DAU/漏斗/LTV)可以继续更新 → 写入 data/官网大盘_渠道.xlsx
  - 国家维度只能冻结在最后一个旧格式日(由 ci_build_payload 自然处理)
  - 策略交叉表 / SKU交叉表 / 引流app 三张表结构未变,仍由 adapt_bi.py 处理

用法(仓库根目录):  BI_FILE=xxx.xlsx python3 ci/adapt_bi_v2.py
"""
import openpyxl, os, re, json, glob, datetime, warnings
warnings.simplefilter("ignore")

D = "ci/data"
OUT = f"{D}/官网大盘_渠道.xlsx"
NEW = os.environ.get("BI_FILE") or sorted(glob.glob("SEO数据看板*.xlsx") + glob.glob("bi_*.xlsx"))[-1]
print("适配输入(新格式):", NEW)

def num(x):
    try: return float(x)
    except Exception: return 0.0
def is2026(v): return bool(re.match(r"2026-\d\d-\d\d", str(v)[:10]))

nb = openpyxl.load_workbook(NEW, read_only=True, data_only=True)
SHEET = next((s for s in ("官网大盘数据", "官网监控明细数据") if s in nb.sheetnames), nb.sheetnames[0])
ws = nb[SHEET]
rows = list(ws.iter_rows(min_row=1, values_only=True))
hdr = [str(x).strip() for x in rows[1]]          # 第2行=字段名
H = {n: i for i, n in enumerate(hdr) if n}
body = [r for r in rows[2:] if r and is2026(r[0])]
if not body:
    raise SystemExit("新格式表里没有 2026 日期行,终止")

def col(*names):
    for n in names:
        if n in H: return H[n]
    raise KeyError("新版大盘表缺列: " + "/".join(names))

# 新旧字段映射(括号内为旧名)
I_DAU   = col("DAU")
I_VIEW  = col("观看uv")
I_REACH = col("快捷支付面板曝光uv", "触达付费集uv")      # 旧:触达付费集uv
I_ORDER = col("商品点击uv", "商品点击/订单创建uv", "创建订单uv")
I_OKR   = col("支付成功率")
I_UV    = col("总付费uv")
I_COIN  = col("金币充值uv")
I_SUB   = col("订阅uv")
I_FO    = col("首订uv")
I_REV   = col("总收入")
I_SUBREV= col("订阅(续订)收入")
I_RET1  = H.get("次日留存率")
I_RN1   = H.get("1期续订人数")
I_FAIL  = H.get("发起支付后支付失败uv")
I_CANC  = H.get("发起支付后取消支付uv")
LTV = [H.get("ltv%d" % k, H.get("LTV%d" % k)) for k in range(31)]

# ---- 按日聚合(跨渠道求和;比率用分子分母重算) ----
acc = {}
for r in body:
    d = str(r[0])[:10]
    a = acc.setdefault(d, {k: 0.0 for k in
        ("dau","view","reach","order","uv","coin","sub","fo","rev","subrev","okw","fail","canc","rn1","retw")})
    a["_ltv"] = a.get("_ltv", [0.0]*31)
    dau = num(r[I_DAU])
    a["dau"]+=dau; a["view"]+=num(r[I_VIEW]); a["reach"]+=num(r[I_REACH]); a["order"]+=num(r[I_ORDER])
    a["uv"]+=num(r[I_UV]); a["coin"]+=num(r[I_COIN]); a["sub"]+=num(r[I_SUB]); a["fo"]+=num(r[I_FO])
    a["rev"]+=num(r[I_REV]); a["subrev"]+=num(r[I_SUBREV])
    a["okw"]+=num(r[I_OKR])*num(r[I_ORDER])                      # 支付成功率按订单加权
    if I_FAIL is not None: a["fail"]+=num(r[I_FAIL])
    if I_CANC is not None: a["canc"]+=num(r[I_CANC])
    if I_RN1  is not None: a["rn1"] +=num(r[I_RN1])
    if I_RET1 is not None: a["retw"]+=num(r[I_RET1])*dau          # 次留 DAU 加权
    for k in range(31):
        if LTV[k] is not None: a["_ltv"][k]+=num(r[LTV[k]])*dau

bad = sorted(d for d, a in acc.items() if a["dau"] <= 0)
if bad:
    print("  [warn] 剔除残缺日(DAU=0,数据未跑完):", bad)
    for d in bad: acc.pop(d)
if not acc:
    raise SystemExit("本次导出无完整日,终止(不写文件)")
cut = min(acc)

# ---- 写出:沿用旧列位,便于 ci_build_payload 复用同一套索引 ----
# 列位: 0日期 1维度 2维度 4DAU 5观看 7触达 9订单 11付费uv 12支付成功率 14金币 16订阅 20续订 22总收入 25订阅收入 30ARPPU 32+LTV
keep = []
if os.path.exists(OUT):
    for r in openpyxl.load_workbook(OUT, read_only=True, data_only=True).worksheets[0].iter_rows(min_row=3, values_only=True):
        if r and is2026(r[0]) and str(r[0])[:10] < cut:
            keep.append(list(r))
wb2 = openpyxl.Workbook(); w = wb2.active; w.title = "官网大盘_渠道"
w.append(["维度"]*3 + ["活跃指标"]*60)
w.append(["日期","口径","备注","", "DAU","观看uv","观看率","快捷支付面板曝光uv","曝光率","商品点击uv","点击率",
          "总付费uv","支付成功率","付费率","金币充值uv","金币率","订阅uv","订阅占比","订阅率","首订uv","续订uv",
          "首订率","总收入","充值收入","金币收入","订阅(续订)收入","首订收入","续订收入","主动充值","ARPU","ARPPU","订阅ARPPU"]
         + ["ltv%d" % k for k in range(31)])
for r in keep: w.append(r)
add = 0
for d in sorted(acc):
    a = acc[d]; dau = a["dau"] or 1
    o = [None]*63
    o[0]=d; o[1]="全渠道合计"; o[2]="SEO+fb+未知"
    o[4]=a["dau"]; o[5]=a["view"]; o[6]=a["view"]/dau
    o[7]=a["reach"]; o[8]=(a["reach"]/a["view"] if a["view"] else 0)
    o[9]=a["order"]; o[10]=(a["order"]/a["reach"] if a["reach"] else 0)
    o[11]=a["uv"]; o[12]=(a["okw"]/a["order"] if a["order"] else 0); o[13]=a["uv"]/dau
    o[14]=a["coin"]; o[16]=a["sub"]; o[19]=a["fo"]; o[20]=max(a["sub"]-a["fo"], 0)
    o[22]=a["rev"]; o[25]=a["subrev"]
    o[29]=a["rev"]/dau; o[30]=(a["rev"]/a["uv"] if a["uv"] else 0)
    for k in range(31): o[32+k]=a["_ltv"][k]/dau
    w.append(o); add += 1
wb2.save(OUT)
print("官网大盘_渠道 (界<%s): 保留%d + 新增%d" % (cut, len(keep), add))

# ---- extras.json:支付失败/取消率、次留、1期续订(只增不减) ----
try:
    p = f"{D}/../extras.json"
    ex = json.load(open(p)) if os.path.exists(p) else {}
    for d, a in acc.items():
        if a["order"] <= 0: continue
        new = {"fail_rate": round(a["fail"]/a["order"]*100, 2),
               "cancel_rate": round(a["canc"]/a["order"]*100, 2),
               "ret1": (round(a["retw"]/a["dau"]*100, 2) if a["retw"] > 0 else None),
               "renew1": (round(a["rn1"]/a["fo"]*100, 2) if (a["fo"] and a["rn1"] > 0) else None)}
        old = ex.get(d) or {}
        for k in ("ret1", "renew1"):
            if new[k] is None: new[k] = old.get(k)
            elif old.get(k) is not None: new[k] = max(old[k], new[k])
        ex[d] = new
    json.dump(dict(sorted(ex.items())), open(p, "w"), ensure_ascii=False, indent=1)
    print("  extras.json 更新 %d 天(累计 %d 天)" % (len(acc), len(ex)))
except Exception as e:
    print("  [warn] extras.json 生成失败:", e)

# ---- 渠道分解:跨导出累积(每份导出只含 2-3 天,不累积会只剩最后一次的) ----
_cp = f"{D}/../channel_daily.json"
ch = json.load(open(_cp)) if os.path.exists(_cp) else {}
_n = 0
for r in body:
    d = str(r[0])[:10]
    if d in bad: continue
    ch.setdefault(d, {})[str(r[1]).strip()] = {
        "dau": num(r[I_DAU]), "rev": num(r[I_REV]), "uv": num(r[I_UV]),
        "view": num(r[I_VIEW]), "order": num(r[I_ORDER]), "sub": num(r[I_SUB]), "fo": num(r[I_FO])}
    _n += 1
json.dump(dict(sorted(ch.items())), open(_cp, "w"), ensure_ascii=False, indent=1)
print("  channel_daily.json:本次写入 %d 行,累计 %d 天" % (_n, len(ch)))
